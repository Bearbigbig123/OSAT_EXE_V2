"""Generate a small, subtle By Tool Shift demonstration fixture.

The demo contains four tools. Tool03 and Tool04 move slightly in opposite
directions during the weekly period. Both trigger only the P50 shift rule,
while all weekly points remain inside the pooled control limits.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from scipy.stats import norm

import oob_module_NGK_nostatic as oob


ROOT = Path(__file__).resolve().parent
DEMO_DIR = ROOT / "test_data" / "by_tool_shift_demo_4tools"
RAW_DIR = DEMO_DIR / "raw_charts"
CSV_NAME = "PK_DEMO_01_SubtleByToolShift_4Tools.csv"
SEED = 20261027
BASELINE_POINTS = 58
WEEKLY_POINTS = 21
CONTROL_SIGMA_MULTIPLIER = 4.0
BASELINE_START = pd.Timestamp("2026-01-01 08:00:00")
WEEKLY_START = pd.Timestamp("2026-06-24 00:00:00")


def standardized_normal(n: int) -> np.ndarray:
    values = norm.ppf((np.arange(n, dtype=float) + 0.5) / n)
    return (values - values.mean()) / values.std(ddof=1)


def write_workbook(path: Path, sheets: list[tuple[str, pd.DataFrame]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, frame in sheets:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(list(frame.columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 42)
            sheet.column_dimensions[cells[0].column_letter].width = width
    workbook.save(path)


def generate() -> None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(SEED)
    baseline_z = standardized_normal(BASELINE_POINTS)
    weekly_z = standardized_normal(WEEKLY_POINTS)
    offsets = (-0.015, -0.005, 0.005, 0.015)
    weekly_shifts = (0.0, 0.0, 0.065, -0.065)
    rows: list[dict] = []

    for tool_index, (offset, shift) in enumerate(zip(offsets, weekly_shifts), 1):
        tool = f"Tool{tool_index:02d}"
        center = 10.0 + offset
        baseline_values = center + 0.040 * baseline_z[rng.permutation(BASELINE_POINTS)]
        weekly_values = center + shift + 0.035 * weekly_z[rng.permutation(WEEKLY_POINTS)]

        for point_index, (point_time, point_val) in enumerate(
            zip(pd.date_range(BASELINE_START, periods=BASELINE_POINTS, freq="3D"), baseline_values), 1
        ):
            rows.append({
                "GroupName": "PK_DEMO",
                "ChartName": "01_SubtleByToolShift_4Tools",
                "point_time": point_time,
                "point_val": round(float(point_val), 5),
                "Batch_ID": f"BSL_{tool}_{point_index:03d}",
                "Matching": tool,
            })
        for point_index, (point_time, point_val) in enumerate(
            zip(pd.date_range(WEEKLY_START, periods=WEEKLY_POINTS, freq="8h"), weekly_values), 1
        ):
            rows.append({
                "GroupName": "PK_DEMO",
                "ChartName": "01_SubtleByToolShift_4Tools",
                "point_time": point_time,
                "point_val": round(float(point_val), 5),
                "Batch_ID": f"WK_{tool}_{point_index:03d}",
                "Matching": tool,
            })

    frame = pd.DataFrame(rows).sort_values(
        ["point_time", "Matching", "Batch_ID"], kind="stable"
    ).reset_index(drop=True)
    baseline = frame[frame["Batch_ID"].str.startswith("BSL_")].copy()
    weekly = frame[frame["Batch_ID"].str.startswith("WK_")].copy()
    center = float(baseline["point_val"].mean())
    sigma = float(baseline["point_val"].std(ddof=1))

    chart = {
        "GroupName": "PK_DEMO",
        "ChartName": "01_SubtleByToolShift_4Tools",
        "ChartID": "PKD001",
        "Material_no": "MAT_PKD001",
        "Target": 10.0,
        "UCL": center + CONTROL_SIGMA_MULTIPLIER * sigma,
        "LCL": center - CONTROL_SIGMA_MULTIPLIER * sigma,
        "USL": 12.0,
        "LSL": 8.0,
        "Characteristics": "Nominal",
        "DetectionLimit": np.nan,
        "ExpectedPattern": "SubtleOppositeShiftWithoutOOC",
        "SampleCount": len(frame),
        "Resolution": 0.01,
    }
    result = oob.by_tool_median_shift_calculator(frame, baseline, weekly, chart)
    records = json.loads(result["by_tool_median_shift_all_tools_json"])
    alarm_records = [record for record in records if record["highlight"]]
    weekly_ooc = int(
        ((weekly["point_val"] > chart["UCL"]) | (weekly["point_val"] < chart["LCL"])).sum()
    )

    if [record["tool"] for record in alarm_records] != ["Tool03", "Tool04"]:
        raise RuntimeError(f"Unexpected alarm tools: {alarm_records}")
    if any(record["triggered_rules"] != ["P50_shift"] for record in alarm_records):
        raise RuntimeError(f"Demo must trigger P50 only: {alarm_records}")
    if weekly_ooc != 0:
        raise RuntimeError(f"Demo must not contain weekly OOC points: {weekly_ooc}")

    csv_frame = frame.copy()
    csv_frame["point_time"] = csv_frame["point_time"].dt.strftime("%Y/%m/%d %H:%M")
    csv_frame.to_csv(RAW_DIR / CSV_NAME, index=False, encoding="utf-8-sig", float_format="%.5f")

    chart_frame = pd.DataFrame([chart])
    summary_frame = pd.DataFrame([{
        "ChartName": chart["ChartName"],
        "HL_by_tool_shift": result["HL_by_tool_shift"],
        "AlarmTools": ", ".join(record["tool"] for record in alarm_records),
        "TriggeredRules": "; ".join(
            f"{record['tool']}:{'/'.join(record['triggered_rules'])}" for record in alarm_records
        ),
        "DisplayedK": "; ".join(
            f"{record['tool']}={record['p50_k']:.2f}" for record in alarm_records
        ),
        "WeeklyOOCPoints": weekly_ooc,
        "Tool01Shift": 0.0,
        "Tool02Shift": 0.0,
        "Tool03Shift": 0.065,
        "Tool04Shift": -0.065,
    }])
    write_workbook(DEMO_DIR / "All_Chart_Information.xlsx", [("Sheet1", chart_frame)])
    write_workbook(
        DEMO_DIR / "OOB_ByTool_Shift_Demo_Workbook.xlsx",
        [("Chart", chart_frame), ("Time", frame), ("Expected_Result", summary_frame)],
    )
    summary_frame.to_csv(DEMO_DIR / "expected_result.csv", index=False, encoding="utf-8-sig")

    preview_chart = dict(chart)
    preview_chart.update({
        "group_name": chart["GroupName"],
        "chart_name": chart["ChartName"],
        "by_tool_median_shift_all_tools_json": result["by_tool_median_shift_all_tools_json"],
    })
    canvas = oob.plot_spc_by_tool_group(
        frame,
        preview_chart,
        WEEKLY_START,
        weekly["point_time"].max(),
        oob_info="HL_by_tool_shift",
        figure_size=(13, 3.2),
    )
    canvas.figure.savefig(DEMO_DIR / "ByTool_Shift_Demo_Preview.png", dpi=160, bbox_inches="tight")

    readme = (
        "By Tool Shift subtle demo\n\n"
        "- 4 tools; 58 baseline points and 21 weekly points per tool (3 points/day for 7 days).\n"
        "- Control limits: pooled baseline center +/- 4 sigma.\n"
        "- Tool01 / Tool02: no shift.\n"
        "- Tool03: +0.065 shift; Tool04: -0.065 shift.\n"
        "- Expected alarm: Tool03 and Tool04, P50 only, displayed K about 2.07.\n"
        "- Expected weekly OOC points: 0.\n"
        "- The opposite shifts largely cancel in the pooled view, demonstrating why By Tool analysis is useful.\n"
    )
    (DEMO_DIR / "README.txt").write_text(readme, encoding="utf-8")

    print(summary_frame.to_string(index=False))
    print(f"Generated demo at: {DEMO_DIR}")


if __name__ == "__main__":
    generate()
