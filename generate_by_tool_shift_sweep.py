"""Rebuild the 25-chart by-tool P05/P50/P95 shift fixture.

The fixture is deterministic.  Baseline control limits are calculated from the
pooled baseline values with sample standard deviation (ddof=1): normally +/-3S,
with the named WideCL and NarrowCL cases using +/-5S and +/-2S respectively.

Run from the project root:

    python generate_by_tool_shift_sweep.py --sync-input
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from scipy.stats import norm

import oob_module_NGK_nostatic as oob


ROOT = Path(__file__).resolve().parent
FIXTURE_DIR = ROOT / "test_data" / "by_tool_median_shift_sweep_25"
RAW_DIR = FIXTURE_DIR / "raw_charts"
INPUT_DIR = ROOT / "input"
INPUT_RAW_DIR = INPUT_DIR / "raw_charts"
SEED = 20261005
BASELINE_POINTS = 58
WEEKLY_POINTS = 7
BASELINE_START = pd.Timestamp("2026-01-01 08:00:00")
WEEKLY_START = pd.Timestamp("2026-06-24 08:00:00")
WIDE_TOOL_SIGMA = 0.16


@dataclass(frozen=True)
class Scenario:
    chart_name: str
    chart_id: str
    material_no: str
    target: float
    usl: float
    lsl: float
    characteristics: str
    detection_limit: float | None
    expected_pattern: str
    resolution: float
    n_tools: int
    shifted_tools: int
    input_shift: float
    weekly_sigma: float
    shifted_sigma: float
    baseline_sigma: float
    baseline_wide_tools: int
    weekly_wide_tools: int
    sparse_weekly: bool
    csv_name: str


def _as_optional_float(value: object) -> float | None:
    return None if pd.isna(value) else float(value)


def load_scenarios() -> list[Scenario]:
    """Read stable scenario inputs from the existing fixture summary."""
    source = FIXTURE_DIR / "expected_hl_summary.csv"
    frame = pd.read_csv(source)
    scenarios: list[Scenario] = []
    for row in frame.to_dict("records"):
        csv_name = Path(str(row["csv_file"]).replace("\\", "/")).name
        scenarios.append(
            Scenario(
                chart_name=str(row["ChartName"]),
                chart_id=str(row["ChartID"]),
                material_no=str(row["Material_no"]),
                target=float(row["Target"]),
                usl=float(row["USL"]),
                lsl=float(row["LSL"]),
                characteristics=str(row["Characteristics"]),
                detection_limit=_as_optional_float(row.get("DetectionLimit")),
                expected_pattern=str(row["ExpectedPattern"]),
                resolution=float(row["Resolution"]),
                n_tools=int(row["n_tools"]),
                shifted_tools=int(row["shifted_tools"]),
                input_shift=float(row["input_shift"]),
                weekly_sigma=float(row["weekly_sigma"]),
                shifted_sigma=float(row["shifted_sigma"]),
                baseline_sigma=float(row["baseline_sigma"]),
                baseline_wide_tools=int(row["baseline_wide_tools"]),
                weekly_wide_tools=int(row["weekly_wide_tools"]),
                sparse_weekly=str(row["sparse_weekly"]).strip().lower() == "true",
                csv_name=csv_name,
            )
        )
    if len(scenarios) != 25:
        raise ValueError(f"Expected 25 scenarios, found {len(scenarios)}")
    return scenarios


def _standard_normal_fixture(n: int, rng: np.random.Generator) -> np.ndarray:
    """Return stratified normal scores with exact sample mean 0 and sample S 1."""
    probabilities = (np.arange(n, dtype=float) + 0.5) / n
    values = norm.ppf(probabilities)
    values = (values - values.mean()) / values.std(ddof=1)
    return values[rng.permutation(n)]


def _tool_offsets(n_tools: int) -> np.ndarray:
    # Small persistent equipment offsets add realism without creating a shift.
    offsets = np.linspace(-0.02, 0.02, n_tools)
    return offsets - offsets.mean()


def _control_multiplier(chart_name: str) -> float:
    if "WideCL" in chart_name:
        return 5.0
    if "NarrowCL" in chart_name:
        return 2.0
    return 3.0


def build_scenario_data(scenario: Scenario, scenario_index: int) -> tuple[pd.DataFrame, dict]:
    rng = np.random.default_rng(SEED + scenario_index * 1009)
    offsets = _tool_offsets(scenario.n_tools)
    shifted_start = scenario.n_tools - scenario.shifted_tools
    baseline_wide = set(range(scenario.baseline_wide_tools))
    weekly_wide = set(range(scenario.n_tools - scenario.weekly_wide_tools, scenario.n_tools))
    weekly_n = 2 if scenario.sparse_weekly else WEEKLY_POINTS
    rows: list[dict] = []

    for tool_idx in range(scenario.n_tools):
        tool = f"Tool{tool_idx + 1:02d}"
        center = scenario.target + offsets[tool_idx]
        baseline_sigma = WIDE_TOOL_SIGMA if tool_idx in baseline_wide else scenario.baseline_sigma
        baseline_values = center + baseline_sigma * _standard_normal_fixture(BASELINE_POINTS, rng)

        is_shifted = tool_idx >= shifted_start and scenario.shifted_tools > 0
        weekly_sigma = scenario.shifted_sigma if is_shifted else scenario.weekly_sigma
        if tool_idx in weekly_wide:
            weekly_sigma = WIDE_TOOL_SIGMA
        weekly_shift = scenario.input_shift if is_shifted else 0.0
        weekly_values = center + weekly_shift + weekly_sigma * _standard_normal_fixture(weekly_n, rng)

        # Five decimals retain the intended sigma while Resolution remains the
        # algorithm's minimum meaningful difference setting.
        baseline_values = np.round(baseline_values, 5)
        weekly_values = np.round(weekly_values, 5)

        baseline_dates = pd.date_range(BASELINE_START, periods=BASELINE_POINTS, freq="3D")
        weekly_dates = pd.date_range(WEEKLY_START + pd.Timedelta(days=WEEKLY_POINTS - weekly_n), periods=weekly_n, freq="D")
        for point_idx, (point_time, point_val) in enumerate(zip(baseline_dates, baseline_values), 1):
            rows.append(
                {
                    "GroupName": "PK_SWEEP",
                    "ChartName": scenario.chart_name,
                    "point_time": point_time,
                    "point_val": float(point_val),
                    "Batch_ID": f"BSL_{tool}_{point_idx:03d}",
                    "Matching": tool,
                }
            )
        for point_idx, (point_time, point_val) in enumerate(zip(weekly_dates, weekly_values), 1):
            rows.append(
                {
                    "GroupName": "PK_SWEEP",
                    "ChartName": scenario.chart_name,
                    "point_time": point_time,
                    "point_val": float(point_val),
                    "Batch_ID": f"WK_{tool}_{point_idx:03d}",
                    "Matching": tool,
                }
            )

    frame = pd.DataFrame(rows).sort_values(["point_time", "Matching", "Batch_ID"], kind="stable").reset_index(drop=True)
    baseline = frame[frame["Batch_ID"].str.startswith("BSL_")]
    center = float(baseline["point_val"].mean())
    baseline_s = float(baseline["point_val"].std(ddof=1))
    multiplier = _control_multiplier(scenario.chart_name)
    control = {
        "Center": center,
        "Baseline_S": baseline_s,
        "Control_Sigma_Multiplier": multiplier,
        "UCL": center + multiplier * baseline_s,
        "LCL": center - multiplier * baseline_s,
    }
    return frame, control


def _chart_row(scenario: Scenario, control: dict, sample_count: int) -> dict:
    return {
        "GroupName": "PK_SWEEP",
        "ChartName": scenario.chart_name,
        "ChartID": scenario.chart_id,
        "Material_no": scenario.material_no,
        "Target": scenario.target,
        "UCL": control["UCL"],
        "LCL": control["LCL"],
        "USL": scenario.usl,
        "LSL": scenario.lsl,
        "Characteristics": scenario.characteristics,
        "DetectionLimit": scenario.detection_limit,
        "ExpectedPattern": scenario.expected_pattern,
        "SampleCount": sample_count,
        "Resolution": scenario.resolution,
    }


def _official_result(frame: pd.DataFrame, chart: dict) -> dict:
    baseline = frame[frame["Batch_ID"].str.startswith("BSL_")].copy()
    weekly = frame[frame["Batch_ID"].str.startswith("WK_")].copy()
    return oob.by_tool_median_shift_calculator(frame, baseline, weekly, chart, min_points=3)


def _count_ooc(values: pd.Series, lcl: float, ucl: float) -> int:
    return int(((values < lcl) | (values > ucl)).sum())


def _expected_row(scenario: Scenario, frame: pd.DataFrame, control: dict, chart: dict, result: dict) -> dict:
    baseline = frame[frame["Batch_ID"].str.startswith("BSL_")]
    weekly = frame[frame["Batch_ID"].str.startswith("WK_")]
    records = json.loads(result.get("by_tool_median_shift_all_tools_json", "[]"))
    triggered = [record for record in records if record.get("highlight")]
    triggered_tools = ",".join(record["tool"] for record in triggered) or "N/A"
    triggered_rules = ";".join(
        f"{record['tool']}:{'/'.join(rule.replace('_shift', '') for rule in record['triggered_rules'])}"
        for record in triggered
    ) or "N/A"
    golden_means = baseline.groupby("Matching")["point_val"].mean()
    golden_tool = str((golden_means - control["Center"]).abs().idxmin())
    return {
        **chart,
        "Center": control["Center"],
        "Baseline_S": control["Baseline_S"],
        "Control_Sigma_Multiplier": control["Control_Sigma_Multiplier"],
        "HL_by_tool_shift": result["HL_by_tool_shift"],
        "triggered_tools": triggered_tools,
        "triggered_rules": triggered_rules,
        "golden_tool": golden_tool,
        "max_tool": result.get("by_tool_median_shift_max_tool", "N/A"),
        "max_diff": result.get("by_tool_median_shift_max_diff", np.nan),
        "max_k": result.get("by_tool_median_shift_max_k", np.nan),
        "eligible_tools": result.get("by_tool_median_shift_tool_count", 0),
        "weekly_ooc_points": _count_ooc(weekly["point_val"], control["LCL"], control["UCL"]),
        "baseline_ooc_points": _count_ooc(baseline["point_val"], control["LCL"], control["UCL"]),
        "weekly_points": len(weekly),
        "baseline_points": len(baseline),
        "n_tools": scenario.n_tools,
        "shifted_tools": scenario.shifted_tools,
        "input_shift": scenario.input_shift,
        "weekly_sigma": scenario.weekly_sigma,
        "shifted_sigma": scenario.shifted_sigma,
        "baseline_sigma": scenario.baseline_sigma,
        "baseline_wide_tools": scenario.baseline_wide_tools,
        "weekly_wide_tools": scenario.weekly_wide_tools,
        "sparse_weekly": scenario.sparse_weekly,
        "csv_file": f"raw_charts\\{scenario.csv_name}",
    }


def _write_styled_workbook(path: Path, sheets: list[tuple[str, pd.DataFrame]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, frame in sheets:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(list(frame.columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 45)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _backup_and_sync_input() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = INPUT_DIR / "backups" / f"by_tool_shift_before_{timestamp}"
    backup_raw = backup / "raw_charts"
    backup_raw.mkdir(parents=True, exist_ok=False)
    workbook = INPUT_DIR / "All_Chart_Information.xlsx"
    if workbook.exists():
        shutil.copy2(workbook, backup / workbook.name)
    for source in INPUT_RAW_DIR.glob("PK_SWEEP_*.csv"):
        shutil.copy2(source, backup_raw / source.name)

    INPUT_RAW_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FIXTURE_DIR / "All_Chart_Information.xlsx", workbook)
    for source in RAW_DIR.glob("PK_SWEEP_*.csv"):
        shutil.copy2(source, INPUT_RAW_DIR / source.name)
    return backup


def run_monte_carlo(trials: int) -> pd.DataFrame:
    """Vectorized reproduction of the current nominal K>2 false-alarm rule.

    Weekly count is seven, so the production calculator takes its ``data_cnt
    >= 2`` path and does not apply rolling confirmation.  The equations below
    mirror its denominator rounding, resolution gate, and nominal direction
    filters while making thousands of trials practical.
    """
    rng = np.random.default_rng(SEED + 900_001)
    rows = []
    for n_tools in (8, 14, 20):
        baseline = rng.normal(10.0, 0.04, size=(trials, n_tools, BASELINE_POINTS))
        weekly = rng.normal(10.0, 0.04, size=(trials, n_tools, WEEKLY_POINTS))
        center = baseline.mean(axis=(1, 2))
        sigma = baseline.reshape(trials, -1).std(axis=1, ddof=1)
        ucl = center + 3 * sigma
        lcl = center - 3 * sigma

        bp05, bp25, bp50, bp75, bp95, bp99865, bp00135 = np.percentile(
            baseline, [5, 25, 50, 75, 95, 99.865, 0.135], axis=2
        )
        wp05, wp50, wp95 = np.percentile(weekly, [5, 50, 95], axis=2)
        p95_denom = np.round(
            np.maximum((bp99865 - bp50) / 3, (ucl[:, None] - bp50) / 6), 8
        )
        p50_denom = np.round(
            np.maximum((bp99865 - bp00135) / 6, (ucl - lcl)[:, None] / 12), 8
        )
        p05_denom = np.round(
            np.maximum((bp50 - bp00135) / 3, (bp50 - lcl[:, None]) / 6), 8
        )

        d95 = wp95 - bp95
        d50 = wp50 - bp50
        d05 = wp05 - bp05
        k95 = np.round(np.abs(d95), 8) / p95_denom
        k50 = np.round(np.abs(d50), 8) / p50_denom
        k05 = np.round(np.abs(d05), 8) / p05_denom
        h95 = (np.abs(d95) >= 0.01) & (k95 > 2) & (d95 > 0)
        h50 = (np.abs(d50) >= 0.01) & (k50 > 2) & ((wp50 < bp25) | (wp50 > bp75))
        h05 = (np.abs(d05) >= 0.01) & (k05 > 2) & (d05 < 0)
        tool_highlight = h95 | h50 | h05
        tool_alerts = int(tool_highlight.sum())
        chart_alerts = int(tool_highlight.any(axis=1).sum())
        total_tools = trials * n_tools
        rows.append(
            {
                "seed": SEED + 900_001,
                "trials": trials,
                "n_tools": n_tools,
                "tool_false_alarm_rate": tool_alerts / total_tools,
                "chart_false_alarm_rate": chart_alerts / trials,
            }
        )
    return pd.DataFrame(rows)


def validate(chart_frame: pd.DataFrame, expected: pd.DataFrame, data_by_name: dict[str, pd.DataFrame]) -> list[str]:
    errors: list[str] = []
    pure_no_shift = {"01_NoShift_8Tools", "06_NoShift_14Tools"}
    for row in expected.to_dict("records"):
        name = row["ChartName"]
        frame = data_by_name[name]
        baseline = frame[frame.Batch_ID.str.startswith("BSL_")]
        multiplier = float(row["Control_Sigma_Multiplier"])
        center = float(baseline.point_val.mean())
        sigma = float(baseline.point_val.std(ddof=1))
        if not np.isclose(row["UCL"], center + multiplier * sigma, atol=1e-12):
            errors.append(f"{name}: invalid UCL")
        if not np.isclose(row["LCL"], center - multiplier * sigma, atol=1e-12):
            errors.append(f"{name}: invalid LCL")
        if name in pure_no_shift and row["HL_by_tool_shift"] != "NO_HIGHLIGHT":
            errors.append(f"{name}: pure NoShift fixture produced a false alert")
        scenario = chart_frame[chart_frame.ChartName == name].iloc[0]
        if int(scenario.SampleCount) != len(frame):
            errors.append(f"{name}: SampleCount mismatch")
    if len(chart_frame) != 25 or len(expected) != 25 or len(data_by_name) != 25:
        errors.append("Fixture must contain exactly 25 charts")
    return errors


def generate(sync_input: bool, monte_carlo_trials: int) -> None:
    scenarios = load_scenarios()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    chart_rows: list[dict] = []
    expected_rows: list[dict] = []
    all_data: list[pd.DataFrame] = []
    data_by_name: dict[str, pd.DataFrame] = {}

    for index, scenario in enumerate(scenarios, 1):
        frame, control = build_scenario_data(scenario, index)
        chart = _chart_row(scenario, control, len(frame))
        result = _official_result(frame, chart)
        expected = _expected_row(scenario, frame, control, chart, result)
        frame_to_write = frame.copy()
        frame_to_write["point_time"] = frame_to_write["point_time"].dt.strftime("%Y/%m/%d %H:%M")
        frame_to_write.to_csv(RAW_DIR / scenario.csv_name, index=False, encoding="utf-8-sig", float_format="%.5f")
        chart_rows.append(chart)
        expected_rows.append(expected)
        all_data.append(frame)
        data_by_name[scenario.chart_name] = frame

    chart_frame = pd.DataFrame(chart_rows)
    expected_frame = pd.DataFrame(expected_rows)
    all_data_frame = pd.concat(all_data, ignore_index=True)
    errors = validate(chart_frame, expected_frame, data_by_name)
    if errors:
        raise RuntimeError("Fixture validation failed:\n- " + "\n- ".join(errors))

    _write_styled_workbook(FIXTURE_DIR / "All_Chart_Information.xlsx", [("Sheet1", chart_frame)])
    expected_frame.to_csv(FIXTURE_DIR / "expected_hl_summary.csv", index=False, encoding="utf-8-sig", float_format="%.10g")
    _write_styled_workbook(FIXTURE_DIR / "expected_hl_summary.xlsx", [("Expected_HL_Summary", expected_frame)])
    _write_styled_workbook(
        FIXTURE_DIR / "OOB_ByTool_MedianShift_Sweep25_Workbook.xlsx",
        [("Chart", chart_frame), ("Time", all_data_frame), ("Expected_HL_Summary", expected_frame)],
    )

    monte_carlo = run_monte_carlo(monte_carlo_trials)
    monte_carlo.to_csv(FIXTURE_DIR / "monte_carlo_false_alarm.csv", index=False, encoding="utf-8-sig", float_format="%.8f")
    files = [FIXTURE_DIR / "All_Chart_Information.xlsx", FIXTURE_DIR / "expected_hl_summary.csv", *sorted(RAW_DIR.glob("PK_SWEEP_*.csv"))]
    manifest = {
        "generator": Path(__file__).name,
        "seed": SEED,
        "control_definition": "pooled baseline mean +/- multiplier * sample S (ddof=1)",
        "default_multiplier": 3,
        "wide_multiplier": 5,
        "narrow_multiplier": 2,
        "files": {str(path.relative_to(FIXTURE_DIR)).replace("\\", "/"): _sha256(path) for path in files},
    }
    (FIXTURE_DIR / "generation_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    backup = None
    if sync_input:
        backup = _backup_and_sync_input()

    print(f"Generated {len(chart_frame)} charts and {len(all_data_frame)} points")
    print(f"Pure NoShift alerts: {expected_frame[expected_frame.ChartName.isin(['01_NoShift_8Tools', '06_NoShift_14Tools'])].HL_by_tool_shift.tolist()}")
    print(monte_carlo.to_string(index=False))
    if backup:
        print(f"Input backup: {backup}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sync-input", action="store_true", help="Back up and replace the active input workbook and PK_SWEEP CSV files")
    parser.add_argument("--monte-carlo-trials", type=int, default=1000, help="No-shift trials for each tool-count case")
    args = parser.parse_args()
    if args.monte_carlo_trials < 1:
        parser.error("--monte-carlo-trials must be >= 1")
    generate(args.sync_input, args.monte_carlo_trials)


if __name__ == "__main__":
    main()
